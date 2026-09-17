"""Lectura y mapeo de columnas de un Excel a filas importables (db.FilaImport)."""
import json
import unicodedata

import openpyxl

from config import get_data_dir
from db import FilaImport
from fechas import formatear_fecha

CAMPOS_DB = ["nombre", "contacto", "estado", "notas", "recomendado_por"]
CAMPOS_OBLIGATORIOS = {"nombre", "contacto"}

MAPEO_PATH = get_data_dir() / "import_mapping.json"

# Palabras clave para adivinar a qué campo corresponde cada header, cuando todavía no hay un
# mapeo guardado (primera vez que se importa un archivo). Se buscan como substring dentro del
# header normalizado (sin tildes, minúscula), en orden: la primera que matchea gana.
PALABRAS_CLAVE = {
    "nombre": ["nombre", "paciente", "cliente", "apellido"],
    "contacto": ["telefono", "contacto", "celular", "movil", "whatsapp", "phone", "tel"],
    "estado": ["estado", "status"],
    "notas": ["nota", "comentario", "observacion"],
    "recomendado_por": ["recomendado", "referido", "derivado"],
}


def _sin_tildes(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c))


def _autodetectar_mapeo(headers: list[str]) -> dict:
    """Heurística simple por palabras clave en el nombre de la columna. Se usa solo para los
    campos que no vinieron del mapeo guardado (o que nunca se importó nada antes)."""
    headers_norm = {h: _sin_tildes(h).lower() for h in headers}
    mapeo = {}
    usados = set()
    for campo, palabras in PALABRAS_CLAVE.items():
        for header, header_norm in headers_norm.items():
            if header in usados:
                continue
            if any(palabra in header_norm for palabra in palabras):
                mapeo[campo] = header
                usados.add(header)
                break
    return mapeo


def leer_headers(path: str) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        hoja = wb.active
        primera_fila = next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [str(h).strip() if h is not None else "" for h in primera_fila]
    finally:
        wb.close()


def cargar_ultimo_mapeo() -> dict:
    """Devuelve {campo_db: header_excel}. Vacío si nunca se importó nada."""
    if not MAPEO_PATH.exists():
        return {}
    try:
        return json.loads(MAPEO_PATH.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}


def guardar_mapeo(mapeo: dict) -> None:
    try:
        MAPEO_PATH.write_text(json.dumps(mapeo, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        pass  # no crashear si no se puede escribir el archivo de config


def sugerir_mapeo(headers: list[str]) -> dict:
    """Primero intenta con el último mapeo guardado (solo para los campos cuyo header sigue
    existiendo en el archivo actual). Para los campos que queden sin cubrir, completa con una
    detección automática por palabras clave en el nombre de columna, así la primera vez que
    se importa un archivo no hay que elegir todo a mano si los headers son razonables
    ("Nombre", "Teléfono", etc.). El usuario siempre puede corregir la sugerencia antes de
    confirmar."""
    guardado = cargar_ultimo_mapeo()
    mapeo = {campo: header for campo, header in guardado.items() if header in headers}

    for campo, header in _autodetectar_mapeo(headers).items():
        if campo not in mapeo:
            mapeo[campo] = header

    return mapeo


def parsear_filas(path: str, mapeo: dict) -> list[FilaImport]:
    """mapeo: {campo_db: header_excel}. Devuelve una FilaImport por cada fila de datos
    (se saltea la fila de encabezados)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        hoja = wb.active
        filas_raw = list(hoja.iter_rows(min_row=1, values_only=True))
    finally:
        wb.close()

    if not filas_raw:
        return []

    headers = [str(h).strip() if h is not None else "" for h in filas_raw[0]]
    idx_por_header = {h: i for i, h in enumerate(headers)}

    def valor(row, campo) -> str:
        header = mapeo.get(campo)
        if not header or header not in idx_por_header:
            return ""
        idx = idx_por_header[header]
        if idx >= len(row) or row[idx] is None:
            return ""
        return str(row[idx]).strip()

    resultado = []
    for n, row in enumerate(filas_raw[1:], start=2):
        if row is None or all(v is None for v in row):
            continue
        resultado.append(
            FilaImport(
                nombre=valor(row, "nombre").upper(),
                contacto=valor(row, "contacto"),
                estado=valor(row, "estado"),
                notas=valor(row, "notas"),
                recomendado_por=valor(row, "recomendado_por"),
                origen=f"Excel fila {n}",
            )
        )
    return resultado


def exportar_excel(path: str, clientes) -> None:
    """Exporta la vista actual (ya filtrada por el llamador) a un .xlsx."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.append(["Nombre", "Teléfono", "Estado", "Fecha de alta", "Fecha de actualización", "Notas"])
    for c in clientes:
        ws.append(
            [
                c.nombre,
                c.contacto,
                c.estado_nombre,
                formatear_fecha(c.fecha_alta),
                formatear_fecha(c.fecha_actualizacion),
                c.notas,
            ]
        )
    wb.save(path)
